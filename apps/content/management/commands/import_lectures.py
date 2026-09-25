"""
Загрузка каталога лекций из таблицы.

    python manage.py import_lectures info_videos.xlsx

Таблица — та, что ведут организаторы: листы «Сезоны» и «Лекции».
Сама она в репозиторий не кладётся, поэтому результат импорта
выгружается в фикстуру и попадает к коллегам уже готовым:

    python manage.py dumpdata content.Topic content.Lecture --indent 2 \\
        -o apps/content/fixtures/lectures.json

--- О разбиении по темам ---

В таблице темы нет, есть только название лекции. Тема определяется по
ключевым словам в названии (TOPIC_RULES). Правила проверяются сверху
вниз, побеждает первое совпавшее — поэтому частные правила стоят выше
общих: «Знакомство с Python. Решение ОДУ» должно попасть в численные
методы, а не в основы языка.

Что не опознано, остаётся без темы и показывается в разделе «Прочее».
Команда печатает такие лекции списком, чтобы их можно было разложить
руками в админке.
"""

import re
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils.text import slugify

from apps.content.models import Lecture, Topic
from apps.seasons.models import Season

# Темы: код, название, раздел, порядок внутри раздела.
TOPICS = [
    ("mechanics", "Механика", Topic.Section.PHYSICS, 10),
    ("celestial", "Небесная механика и баллистика", Topic.Section.PHYSICS, 20),
    ("fluids", "Гидро- и аэродинамика", Topic.Section.PHYSICS, 30),
    ("thermo", "Термодинамика и атмосфера", Topic.Section.PHYSICS, 40),
    ("electricity", "Электричество и магнетизм", Topic.Section.PHYSICS, 50),
    ("optics", "Оптика и излучение", Topic.Section.PHYSICS, 60),
    ("math", "Математический аппарат", Topic.Section.PHYSICS, 70),
    ("python", "Основы Python", Topic.Section.PROGRAMMING, 10),
    ("numerical", "Численные методы", Topic.Section.PROGRAMMING, 20),
    ("algorithms", "Алгоритмы", Topic.Section.PROGRAMMING, 30),
    ("data", "Данные, статистика и машинное обучение", Topic.Section.PROGRAMMING, 40),
    ("solutions", "Разборы задач", Topic.Section.COMMON, 10),
    ("about", "Об олимпиаде", Topic.Section.COMMON, 20),
]

# Порядок важен: сверху частные случаи, снизу общие.
TOPIC_RULES = [
    ("solutions", ["разбор"]),
    ("about", ["юбиле"]),
    ("numerical", ["решение оду", "интерполляц", "численны", "вычислительная"]),
    ("data", ["статистик", "машинное обучение"]),
    ("algorithms", ["бинарный поиск", "сортировк", "алгоритм", "структуры данных"]),
    ("python", ["python", "numpy", "matplotlib"]),
    ("celestial", ["небесн", "кеплер", "ракет", "мещерск", "реактивн",
                   "солнечный парус", "орбит", "баллистик"]),
    ("fluids", ["бернулли", "эйлер", "аэродинамик", "ударные волны",
                "поверхностное натяжение", "сопротивление движению", "жидкост"]),
    ("optics", ["оптик", "фотометри", "звездны", "рефракц", "излучени",
                "стефана", "освещенн", "спектр", "резонанс"]),
    ("electricity", ["электрич", "магнит", "индукцион", "гаусса", "rc-цеп",
                     "плазма", "магнетизм"]),
    # Механика выше термодинамики: лекция «Вращательное движение. Закон
    # Гука. Тепловое расширение» — про механику, хотя и упоминает нагрев.
    ("mechanics", ["механик", "движени", "колебани", "система отсчета",
                   "системы отсчета", "гука", "вращательн", "зси", "зсэ"]),
    ("thermo", ["барометрическ", "тепловое расширение", "термодинамик", "атмосфер"]),
    ("math", ["производная", "интеграл. вектор", "вектор-функц"]),
]

SECTION_SHEET = "Сезоны"
LECTURE_SHEET = "Лекции"


class Command(BaseCommand):
    help = "Загружает каталог лекций из таблицы организаторов"

    def add_arguments(self, parser):
        parser.add_argument("path", help="Файл .xlsx с листами «Сезоны» и «Лекции»")
        parser.add_argument("--dry-run", action="store_true",
                            help="показать, что получится, не трогая базу")

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError as exc:
            raise CommandError(
                "Нужна библиотека openpyxl: .venv/bin/pip install openpyxl"
            ) from exc

        path = Path(options["path"])
        if not path.is_file():
            raise CommandError(f"Нет такого файла: {path}")

        workbook = openpyxl.load_workbook(path, read_only=True)
        if LECTURE_SHEET not in workbook.sheetnames:
            raise CommandError(f"В файле нет листа «{LECTURE_SHEET}»")

        rows = list(self._read(workbook[LECTURE_SHEET]))
        if options["dry_run"]:
            self._report(rows)
            return

        with transaction.atomic():
            topics = self._ensure_topics()
            created, updated, unmatched = self._save(rows, topics)

        self.stdout.write(self.style.SUCCESS(
            f"Лекций создано: {created}, обновлено: {updated}"
        ))
        if unmatched:
            self.stdout.write("Без темы — разложите в админке (Контент → Лекции):")
            for title in unmatched:
                self.stdout.write(f"    {title}")

    def _read(self, sheet):
        """Строки таблицы. Сезон и тур указаны только в первой строке блока."""
        season = stage = None
        for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if number == 1:
                continue  # заголовок
            # В режиме read_only строки бывают короче пяти колонок —
            # дополняем, чтобы не ловить IndexError на пустых хвостах.
            values = [str(c).strip() if c is not None else "" for c in row[:5]]
            values += [""] * (5 - len(values))
            season = values[0] or season
            stage = values[1] or stage
            title, lecturer, url = values[2], values[3], values[4]
            if not title or not url:
                continue
            yield {"season": season, "stage": stage, "title": title,
                   "lecturer": lecturer, "url": url}

    def _ensure_topics(self):
        topics = {}
        for slug, title, section, order in TOPICS:
            topics[slug], _ = Topic.objects.update_or_create(
                slug=slug,
                defaults={"title": title, "section": section, "order": order},
            )
        return topics

    def _save(self, rows, topics):
        created = updated = 0
        unmatched = []
        for position, row in enumerate(rows):
            season = self._season(row["season"])
            if season is None:
                continue

            slug = self._slug(season, row["title"], position)
            topic_slug = match_topic(row["title"])
            if topic_slug is None:
                unmatched.append(row["title"])

            _, was_created = Lecture.objects.update_or_create(
                season=season, slug=slug,
                defaults={
                    "title": row["title"],
                    "lecturer": row["lecturer"],
                    "video_url": row["url"],
                    "topic": topics.get(topic_slug),
                    "description": self._description(row),
                    # Каталог заводят сами организаторы, одобрение не нужно.
                    "status": Lecture.Status.APPROVED,
                },
            )
            created, updated = (created + 1, updated) if was_created else (created, updated + 1)
        return created, updated, unmatched

    def _description(self, row):
        """«отборочный» → «Лекция отборочного тура.»"""
        tour = (row["stage"] or "").strip().lower()
        if not tour or tour == "-":
            return ""
        # В таблице тур записан в именительном падеже: «отборочный», «финальный».
        tour = re.sub(r"(ый|ой)$", "ого", tour)
        tour = re.sub(r"ий$", "его", tour)
        return f"Лекция {tour} тура."

    def _season(self, label):
        """«2024/25» → сезон с годом окончания 2025."""
        match = re.match(r"(\d{4})/(\d{2})", label or "")
        if not match:
            return None
        year = int(match.group(1)) + 1
        season = Season.objects.filter(year=year).first()
        if season:
            return season
        from apps.content.management.commands.import_archive import roman

        return Season.objects.create(
            year=year,
            slug=f"ao{str(year)[-2:]}",
            title=f"Аэрокосмическая олимпиада МФТИ {roman(year)}",
            subtitle=f"Сезон {year - 1}/{str(year)[-2:]}",
            is_published=True,
        )

    def _slug(self, season, title, position):
        """Адрес лекции. Названия на кириллице slugify обнуляет, поэтому
        для них берём порядковый номер — адрес всё равно служебный."""
        base = slugify(title) or f"lecture-{position + 1}"
        slug, number = base, 1
        while Lecture.objects.filter(season=season, slug=slug).exclude(title=title).exists():
            number += 1
            slug = f"{base}-{number}"
        return slug

    def _report(self, rows):
        by_topic = {}
        for row in rows:
            by_topic.setdefault(match_topic(row["title"]) or "— без темы —", []).append(row["title"])
        for slug, titles in sorted(by_topic.items()):
            self.stdout.write(f"{slug} ({len(titles)}):")
            for title in titles:
                self.stdout.write(f"    {title}")


def match_topic(title: str):
    """Тема по названию лекции или None, если не опознали."""
    lowered = title.lower().replace("ё", "е")
    for slug, keywords in TOPIC_RULES:
        if any(word in lowered for word in keywords):
            return slug
    return None
