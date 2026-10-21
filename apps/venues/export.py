"""Выгрузка участников площадки в Excel."""

import io
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def format_passport_data(profile):
    """Строка паспортных данных участника для таблицы.

    Например: «Паспорт РФ: серия 4510 № 123456, выдан 01.02.2020, Отделом УФМС, код подразделения 770-001».
    Корректно обрабатывает пустые или частично заполненные поля.
    """
    if not profile:
        return ""

    label = profile.get_doc_type_display() if profile.doc_type else ""
    parts = []

    series = (profile.doc_series or "").strip()
    number = (profile.doc_number or "").strip()
    if series and number:
        parts.append(f"серия {series} № {number}")
    elif number:
        parts.append(f"№ {number}")
    elif series:
        parts.append(f"серия {series}")

    issued_parts = []
    verb = "выдано" if profile.doc_type == "birth_cert" else "выдан"
    issued_date = (
        profile.doc_issued_at.strftime("%d.%m.%Y") if isinstance(profile.doc_issued_at, date) else ""
    )
    issued_by = (profile.doc_issued_by or "").strip()

    if issued_date and issued_by:
        issued_parts.append(f"{verb} {issued_date}, {issued_by}")
    elif issued_date:
        issued_parts.append(f"{verb} {issued_date}")
    elif issued_by:
        issued_parts.append(f"{verb} {issued_by}")

    if issued_parts:
        parts.append(", ".join(issued_parts))

    division_code = (profile.doc_division_code or "").strip()
    if division_code:
        parts.append(f"код подразделения {division_code}")

    details = ", ".join(parts)
    if label and details:
        return f"{label}: {details}"
    if label:
        return label
    return details


def generate_venue_participants_excel(venue, bookings):
    """Сформировать Excel-книгу (.xlsx) со списком участников площадки."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Участники"

    headers = ["ФИО", "Паспортные данные", "Школа", "Адрес точки проведения"]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True)
    header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    header_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for booking in bookings:
        user = booking.registration.user
        profile = getattr(user, "profile", None)
        full_name = profile.full_name if (profile and profile.full_name) else (user.get_full_name() or user.email)
        passport = format_passport_data(profile)
        school = (profile.school or "").strip() if profile else ""
        address = venue.address or ""
        ws.append([full_name, passport, school, address])

    data_alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 2), min_col=1, max_col=4):
        for cell in row:
            cell.alignment = data_alignment

    min_widths = [30, 45, 30, 35]
    for col_idx, min_w in enumerate(min_widths, 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, ws.max_row + 1))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, min_w), 70)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def generate_offline_grading_excel(stage, rows, problems, request=None):
    """Сформировать Excel-книгу (.xlsx) с ведомостью очного тура:
    ФИО, точка проведения, город, дата, комплект, задачи, сумма, скан работы.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ведомость очного тура"

    headers = ["ФИО участника", "Точка проведения", "Город", "Дата проведения", "Комплект"]
    for p in problems:
        headers.append(f"Задача {p.number}")
    headers.extend(["Сумма", "Скан работы"])

    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True)
    header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    ws.cell(row=1, column=1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    inactive_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
    inactive_font = Font(name="Calibri", size=11, color="999999")
    bold_font = Font(name="Calibri", size=11, bold=True)
    link_font = Font(name="Calibri", size=11, color="0000EE", underline="single")

    for row_idx, r in enumerate(rows, start=2):
        booking = r["booking"]
        user = booking.registration.user
        profile = getattr(user, "profile", None)
        full_name = profile.full_name if (profile and profile.full_name) else (user.get_full_name() or user.email)
        venue_name = booking.venue.short_title or booking.venue.title
        city = booking.venue.city
        held_at = r.get("held_at_val") or ""
        variant = r.get("variant") or ""

        row_data = [full_name, venue_name, city, held_at, variant]

        cells = r.get("cells", [])
        for c in cells:
            if c.get("is_inactive"):
                row_data.append("—")
            else:
                val = c.get("value")
                if val:
                    try:
                        num_val = float(str(val).replace(",", "."))
                        row_data.append(int(num_val) if num_val.is_integer() else num_val)
                    except ValueError:
                        row_data.append(val)
                else:
                    row_data.append("")

        total_val = r.get("total")
        if total_val:
            try:
                num_total = float(str(total_val).replace(",", "."))
                row_data.append(int(num_total) if num_total.is_integer() else num_total)
            except ValueError:
                row_data.append(total_val)
        else:
            row_data.append(0)

        scan_url = r.get("work_scan_url")
        if scan_url:
            full_scan_url = request.build_absolute_uri(scan_url) if request else scan_url
            row_data.append(full_scan_url)
        else:
            row_data.append("Нет скана")

        ws.append(row_data)

        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=row_idx, column=2).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=row_idx, column=3).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_idx, column=4).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_idx, column=5).alignment = Alignment(horizontal="center", vertical="center")

        prob_start_col = 6
        for p_idx, c in enumerate(cells):
            col = prob_start_col + p_idx
            cell = ws.cell(row=row_idx, column=col)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if c.get("is_inactive"):
                cell.fill = inactive_fill
                cell.font = inactive_font

        total_col = prob_start_col + len(cells)
        total_cell = ws.cell(row=row_idx, column=total_col)
        total_cell.alignment = Alignment(horizontal="center", vertical="center")
        total_cell.font = bold_font

        scan_col = total_col + 1
        scan_cell = ws.cell(row=row_idx, column=scan_col)
        scan_cell.alignment = Alignment(horizontal="center", vertical="center")
        if scan_url:
            full_scan_url = request.build_absolute_uri(scan_url) if request else scan_url
            scan_cell.value = "Открыть скан"
            scan_cell.hyperlink = full_scan_url
            scan_cell.font = link_font

    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, ws.max_row + 1))
        min_w = 12
        if col_idx == 1:
            min_w = 28
        elif col_idx == 2:
            min_w = 20
        elif col_idx == 3:
            min_w = 16
        elif col_idx == 4:
            min_w = 16
        elif col_idx == len(headers):
            min_w = 16
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, min_w), 50)

    ws.freeze_panes = "B2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

