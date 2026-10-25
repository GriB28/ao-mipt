"""Тесты площадок: выгрузка данных участников в Excel, права доступа."""

import io
from datetime import date, timedelta
from decimal import Decimal

import openpyxl
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import ConsentDocument, DocumentType, ParticipantProfile, User
from apps.participation.models import Registration, VenueBooking
from apps.seasons.models import Season, Stage
from apps.venues.export import format_passport_data
from apps.venues.models import Venue


class VenueExportUnitTests(TestCase):
    """Модульные тесты форматирования паспортных данных и генерации файла Excel."""

    def test_format_passport_data_empty_or_none(self):
        self.assertEqual(format_passport_data(None), "")
        profile = ParticipantProfile()
        self.assertEqual(format_passport_data(profile), "")

    def test_format_passport_data_passport_rf_complete(self):
        user = User.objects.create_user("u1@b.ru", "pass123")
        profile = ParticipantProfile.objects.create(
            user=user,
            doc_type=DocumentType.PASSPORT_RF,
            doc_series="4510",
            doc_number="123456",
            doc_issued_at=date(2020, 2, 1),
            doc_issued_by="ГУ МВД по г. Москве",
            doc_division_code="770-001",
        )
        expected = "Паспорт РФ: серия 4510 № 123456, выдан 01.02.2020, ГУ МВД по г. Москве, код подразделения 770-001"
        self.assertEqual(format_passport_data(profile), expected)

    def test_format_passport_data_birth_cert(self):
        user = User.objects.create_user("u2@b.ru", "pass123")
        profile = ParticipantProfile.objects.create(
            user=user,
            doc_type=DocumentType.BIRTH_CERT,
            doc_series="IV-МЮ",
            doc_number="654321",
            doc_issued_at=date(2014, 5, 15),
            doc_issued_by="Орган ЗАГС г. Москвы",
        )
        expected = "Свидетельство о рождении: серия IV-МЮ № 654321, выдано 15.05.2014, Орган ЗАГС г. Москвы"
        self.assertEqual(format_passport_data(profile), expected)

    def test_format_passport_data_foreign(self):
        user = User.objects.create_user("u3@b.ru", "pass123")
        profile = ParticipantProfile.objects.create(
            user=user,
            doc_type=DocumentType.FOREIGN,
            doc_number="AB987654",
            doc_issued_at=date(2022, 1, 10),
            doc_issued_by="МВД РК",
        )
        expected = "Документ другой страны: № AB987654, выдан 10.01.2022, МВД РК"
        self.assertEqual(format_passport_data(profile), expected)

    def test_format_passport_data_partial(self):
        user = User.objects.create_user("u4@b.ru", "pass123")
        profile = ParticipantProfile.objects.create(
            user=user,
            doc_type=DocumentType.PASSPORT_RF,
            doc_series="4510",
            doc_number="123456",
        )
        self.assertEqual(format_passport_data(profile), "Паспорт РФ: серия 4510 № 123456")


class VenueExportViewTests(TestCase):
    """Интеграционные тесты выгрузки участников площадки в Excel и прав доступа."""

    def setUp(self):
        now = timezone.now()
        self.season = Season.objects.create(
            year=2099, slug="s2099", title="Сезон 2099", is_active=True, is_published=True
        )
        self.stage = Stage.objects.create(
            season=self.season,
            kind=Stage.Kind.OFFLINE,
            slug="offline-final",
            title="Очный тур",
            starts_at=now + timedelta(days=5),
            ends_at=now + timedelta(days=5, hours=4),
            registration_opens_at=now - timedelta(days=2),
            registration_closes_at=now + timedelta(days=3),
            is_published=True,
        )
        self.organizer = User.objects.create_user("org@b.ru", "pass12345", role=User.Role.ORGANIZER)
        self.other_organizer = User.objects.create_user(
            "other_org@b.ru", "pass12345", role=User.Role.ORGANIZER
        )
        self.admin = User.objects.create_user("admin@b.ru", "pass12345", role=User.Role.ADMIN, is_staff=True)
        self.participant = User.objects.create_user("part@b.ru", "pass12345", role=User.Role.PARTICIPANT)

        self.venue = Venue.objects.create(
            title="Лицей № 1580",
            region="Москва",
            city="Москва",
            address="ул. Талалихина, 1к1",
            latitude=55.73,
            longitude=37.67,
            status=Venue.Status.APPROVED,
        )
        self.venue.managers.add(self.organizer)
        self.venue.seasons.add(self.season)

        # Участник 1 с заполненным профилем
        self.user1 = User.objects.create_user("ivanov@b.ru", "pass12345")
        ParticipantProfile.objects.create(
            user=self.user1,
            last_name="Иванов",
            first_name="Иван",
            middle_name="Иванович",
            school="ГБОУ Лицей № 1502",
            doc_type=DocumentType.PASSPORT_RF,
            doc_series="4510",
            doc_number="123456",
            doc_issued_at=date(2020, 1, 15),
            doc_issued_by="ГУ МВД г. Москвы",
            doc_division_code="770-001",
        )
        reg1 = Registration.objects.create(user=self.user1, season=self.season)
        VenueBooking.objects.create(
            registration=reg1, stage=self.stage, venue=self.venue, status=VenueBooking.Status.BOOKED
        )

        # Участник 2 без профиля
        self.user2 = User.objects.create_user(
            "petrov@b.ru", "pass12345", first_name="Пётр", last_name="Петров"
        )
        reg2 = Registration.objects.create(user=self.user2, season=self.season)
        VenueBooking.objects.create(
            registration=reg2, stage=self.stage, venue=self.venue, status=VenueBooking.Status.BOOKED
        )

        # Отменённая запись (не должна попасть в выгрузку)
        self.user_cancelled = User.objects.create_user("canc@b.ru", "pass12345")
        reg_canc = Registration.objects.create(user=self.user_cancelled, season=self.season)
        VenueBooking.objects.create(
            registration=reg_canc, stage=self.stage, venue=self.venue, status=VenueBooking.Status.CANCELLED
        )

    def test_anonymous_redirected_to_login(self):
        url = reverse("venues:export", args=[self.venue.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 302)
        self.assertIn("/accounts/login/", response["Location"])

    def test_participant_forbidden(self):
        self.client.force_login(self.participant)
        url = reverse("venues:export", args=[self.venue.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 403)

    def test_other_organizer_forbidden(self):
        self.client.force_login(self.other_organizer)
        url = reverse("venues:export", args=[self.venue.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 403)

    def test_venue_manager_can_export(self):
        self.client.force_login(self.organizer)
        url = reverse("venues:export", args=[self.venue.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(f'filename="venue-{self.venue.pk}-participants.xlsx"', response["Content-Disposition"])

        # Проверяем содержимое Excel
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.title, "Участники")

        # Заголовки
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(headers, ["ФИО", "Паспортные данные", "Школа", "Адрес точки проведения"])

        # Строки с данными
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        # Ожидаем 2 записи (отменённая не попадает)
        self.assertEqual(len(rows), 2)

        # Проверяем данные участника 1
        row_ivanov = next(r for r in rows if "Иванов Иван Иванович" in r[0])
        self.assertIn("Паспорт РФ: серия 4510 № 123456", row_ivanov[1])
        self.assertIn("код подразделения 770-001", row_ivanov[1])
        self.assertEqual(row_ivanov[2], "ГБОУ Лицей № 1502")
        self.assertEqual(row_ivanov[3], self.venue.address)

        # Проверяем участника 2 (без профиля ПД и школы)
        row_petrov = next(r for r in rows if "Петров" in r[0] or "petrov@b.ru" in r[0])
        self.assertIn(row_petrov[1], ("", None))
        self.assertIn(row_petrov[2], ("", None))
        self.assertEqual(row_petrov[3], self.venue.address)

        # Участника с отменённой записью быть не должно
        self.assertFalse(any("canc@b.ru" in str(r[0]) for r in rows))

    def test_admin_can_export(self):
        self.client.force_login(self.admin)
        url = reverse("venues:export", args=[self.venue.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)

    def test_button_present_on_detail_page_for_manager(self):
        self.client.force_login(self.organizer)
        url = reverse("venues:detail", args=[self.venue.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, reverse("venues:export", args=[self.venue.pk]))
        self.assertContains(response, "Выгрузить в Excel")

    def test_button_not_present_for_participant(self):
        self.client.force_login(self.participant)
        url = reverse("venues:detail", args=[self.venue.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, reverse("venues:export", args=[self.venue.pk]))


class OfflineGradingAndScansTests(TestCase):
    """Тесты ведомости очного тура, переключения комплектов и загрузки сканов."""

    def setUp(self):
        from decimal import Decimal
        from django.core.files.uploadedfile import SimpleUploadedFile
        from apps.contest.models import Problem, Score

        self.Decimal = Decimal
        self.SimpleUploadedFile = SimpleUploadedFile
        self.Problem = Problem
        self.Score = Score

        now = timezone.now()
        self.season = Season.objects.create(
            year=2099, slug="s2099", title="Сезон 2099", is_active=True, is_published=True
        )
        self.stage = Stage.objects.create(
            season=self.season,
            kind=Stage.Kind.OFFLINE,
            slug="offline-final",
            title="Очный тур",
            starts_at=now - timedelta(days=1),
            ends_at=now + timedelta(days=2),
            is_published=True,
            order=1,
        )
        self.venue = Venue.objects.create(
            title="МФТИ КПМ",
            city="Долгопрудный",
            address="Первомайская 7",
            latitude=55.93,
            longitude=37.52,
            status=Venue.Status.APPROVED,
            default_variant=1,
            event_date=date(2099, 10, 15),
        )
        self.venue.seasons.add(self.season)

        self.organizer = User.objects.create_user("org_offline@b.ru", "pass12345", role=User.Role.ORGANIZER)
        self.participant1 = User.objects.create_user("part1@b.ru", "pass12345", role=User.Role.PARTICIPANT)
        self.participant2 = User.objects.create_user("part2@b.ru", "pass12345", role=User.Role.PARTICIPANT)

        reg1 = Registration.objects.create(user=self.participant1, season=self.season)
        self.booking1 = VenueBooking.objects.create(
            registration=reg1,
            stage=self.stage,
            venue=self.venue,
            status=VenueBooking.Status.CONFIRMED,
        )

        reg2 = Registration.objects.create(user=self.participant2, season=self.season)
        self.booking2 = VenueBooking.objects.create(
            registration=reg2,
            stage=self.stage,
            venue=self.venue,
            status=VenueBooking.Status.CONFIRMED,
        )

        # Задачи: 1.1, 1.2 (Комплект 1), 2.1, 2.2 (Комплект 2)
        self.p1_1 = Problem.objects.create(
            stage=self.stage, number="1.1", title="Задача 1.1",
            max_score=Decimal("10"), status=Problem.Status.APPROVED,
        )
        self.p1_2 = Problem.objects.create(
            stage=self.stage, number="1.2", title="Задача 1.2",
            max_score=Decimal("10"), status=Problem.Status.APPROVED,
        )
        self.p2_1 = Problem.objects.create(
            stage=self.stage, number="2.1", title="Задача 2.1",
            max_score=Decimal("10"), status=Problem.Status.APPROVED,
        )
        self.p2_2 = Problem.objects.create(
            stage=self.stage, number="2.2", title="Задача 2.2",
            max_score=Decimal("10"), status=Problem.Status.APPROVED,
        )

    def test_manager_sees_tabs_participant_does_not(self):
        # Обычный участник: нет вкладок
        self.client.force_login(self.participant1)
        resp = self.client.get(reverse("venues:map"))
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, "sub-tabs-nav")
        self.assertNotContains(resp, "Проверка задач")
        self.assertNotContains(resp, "Загрузка скана работы")

        # Организатор: видит вкладки
        self.client.force_login(self.organizer)
        resp = self.client.get(reverse("venues:map"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "sub-tabs-nav")
        self.assertContains(resp, "Проверка задач")
        self.assertContains(resp, "Загрузка скана работы")

    def test_offline_grading_ajax_scores_and_variant_deactivation(self):
        self.client.force_login(self.organizer)
        save_url = reverse("venues:save_scores")

        # Ставим балл 8.5 за 1.1 (Комплект 1)
        resp = self.client.post(
            save_url,
            {"booking_id": self.booking1.pk, "problem_id": self.p1_1.pk, "points": "8.5"},
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data["status"], "ok")
        self.assertEqual(data["total"], "8.5")

        # Ставим балл 7.0 за 2.1 (Комплект 2) — при активном комплекте 1 он должен быть skipped
        resp = self.client.post(
            save_url,
            {"booking_id": self.booking1.pk, "problem_id": self.p2_1.pk, "points": "7"},
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data["total"], "8.5")  # 7 не входит в сумму!

        s1_1 = self.Score.objects.get(registration=self.booking1.registration, problem=self.p1_1)
        self.assertFalse(s1_1.is_skipped)
        s2_1 = self.Score.objects.get(registration=self.booking1.registration, problem=self.p2_1)
        self.assertTrue(s2_1.is_skipped)

        # Переключаем комплект участника на 2
        resp = self.client.post(
            save_url,
            {"booking_id": self.booking1.pk, "variant": "2"},
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data["variant"], 2)
        self.assertEqual(data["total"], "7")  # теперь в сумме только 2.1!

        s1_1.refresh_from_db()
        s2_1.refresh_from_db()
        self.assertTrue(s1_1.is_skipped)
        self.assertFalse(s2_1.is_skipped)

    def test_offline_grading_score_max_validation(self):
        self.client.force_login(self.organizer)
        save_url = reverse("venues:save_scores")

        # 1. Попытка поставить балл больше максимального (10)
        resp = self.client.post(
            save_url,
            {"booking_id": self.booking1.pk, "problem_id": self.p1_1.pk, "points": "15"},
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 400)
        self.assertIn("error", resp.json())
        self.assertIn("10", resp.json()["error"])
        self.assertFalse(self.Score.objects.filter(registration=self.booking1.registration, problem=self.p1_1).exists())

        # 2. Попытка поставить отрицательный балл
        resp = self.client.post(
            save_url,
            {"booking_id": self.booking1.pk, "problem_id": self.p1_1.pk, "points": "-1"},
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 400)
        self.assertIn("error", resp.json())
        self.assertFalse(self.Score.objects.filter(registration=self.booking1.registration, problem=self.p1_1).exists())

        # 3. Валидный балл (ровно максимум 10)
        resp = self.client.post(
            save_url,
            {"booking_id": self.booking1.pk, "problem_id": self.p1_1.pk, "points": "10"},
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 200)
        score = self.Score.objects.get(registration=self.booking1.registration, problem=self.p1_1)
        self.assertEqual(score.points, Decimal("10"))

        # 4. Проблема без max_score (None) принимает любые неотрицательные баллы
        self.p1_1.max_score = None
        self.p1_1.save()
        resp = self.client.post(
            save_url,
            {"booking_id": self.booking1.pk, "problem_id": self.p1_1.pk, "points": "25"},
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 200)
        score.refresh_from_db()
        self.assertEqual(score.points, Decimal("25"))

        # 5. Проверка отображения max_score в шаблоне ведомости
        self.p1_1.max_score = Decimal("10")
        self.p1_1.save()
        map_resp = self.client.get(f"{reverse('venues:map')}?tab=grading")
        self.assertEqual(map_resp.status_code, 200)
        self.assertContains(map_resp, 'data-max-score="10"')
        self.assertContains(map_resp, "макс: 10")

    def test_offline_scan_upload_download_and_permissions(self):
        self.client.force_login(self.organizer)
        upload_url = reverse("venues:upload_work_scan")

        # Не-PDF отклоняется
        txt_file = self.SimpleUploadedFile("work.txt", b"not a pdf", content_type="text/plain")
        resp = self.client.post(upload_url, {"booking_id": self.booking1.pk, "work_scan": txt_file})
        self.booking1.refresh_from_db()
        self.assertFalse(bool(self.booking1.work_scan))

        # Валидный PDF загружается
        pdf_file = self.SimpleUploadedFile("work.pdf", b"%PDF-1.4 sample content", content_type="application/pdf")
        resp = self.client.post(
            upload_url,
            {"booking_id": self.booking1.pk, "work_scan": pdf_file, "variant": "2", "held_at": "2099-10-16"},
        )
        self.booking1.refresh_from_db()
        self.assertTrue(bool(self.booking1.work_scan))
        self.assertEqual(self.booking1.variant, 2)
        self.assertEqual(str(self.booking1.held_at), "2099-10-16")

        scan_url = self.booking1.get_scan_url()

        # Организатор может скачать/просмотреть скан
        resp = self.client.get(scan_url)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "application/pdf")

        # Участник до публикации результатов получает 404
        self.client.force_login(self.participant1)
        resp = self.client.get(scan_url)
        self.assertEqual(resp.status_code, 404)

        # Публикуем результаты
        self.stage.results_published = True
        self.stage.save()

        # Теперь участник видит свой скан
        resp = self.client.get(scan_url)
        self.assertEqual(resp.status_code, 200)

        # Чужой участник не видит
        self.client.force_login(self.participant2)
        resp = self.client.get(scan_url)
        self.assertEqual(resp.status_code, 404)

        # Удаление скана организатором
        self.client.force_login(self.organizer)
        del_url = reverse("venues:delete_work_scan", args=[self.booking1.pk])
        resp = self.client.post(del_url)
        self.booking1.refresh_from_db()
        self.assertFalse(bool(self.booking1.work_scan))

    def test_export_grading_excel(self):
        self.client.force_login(self.organizer)
        url = reverse("venues:export_grading")
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("vedomost_ochny_tur_", resp["Content-Disposition"])

        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertIn("Общая ведомость", wb.sheetnames)
        self.assertIn("Долгопрудный", wb.sheetnames)

        ws_all = wb["Общая ведомость"]
        self.assertEqual(ws_all.cell(row=1, column=1).value, "ФИО участника")
        names_all = [ws_all.cell(row=r, column=1).value for r in range(2, ws_all.max_row + 1)]
        self.assertIn("part1@b.ru", names_all)
        self.assertIn("part2@b.ru", names_all)

        ws_city = wb["Долгопрудный"]
        self.assertEqual(ws_city.cell(row=1, column=1).value, "ФИО участника")
        names_city = [ws_city.cell(row=r, column=1).value for r in range(2, ws_city.max_row + 1)]
        self.assertIn("part1@b.ru", names_city)
        self.assertIn("part2@b.ru", names_city)

    def test_participant_views_results_and_scan_on_offline_page(self):
        # 1. До публикации результатов участник видит только статус записи, но не баллы и не результаты
        self.stage.results_published = False
        self.stage.save()

        self.Score.objects.create(
            registration=self.booking1.registration,
            problem=self.p1_1,
            points=self.Decimal("8"),
            venue=self.venue,
        )
        self.Score.objects.create(
            registration=self.booking1.registration,
            problem=self.p1_2,
            points=self.Decimal("10"),
            venue=self.venue,
        )

        self.client.force_login(self.participant1)
        resp = self.client.get(reverse("venues:map"))
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, "Ваши результаты очного тура")
        self.assertNotContains(resp, "Итоговый балл:")

        # 2. Публикуем результаты этапа
        self.stage.results_published = True
        self.stage.save()

        resp = self.client.get(reverse("venues:map"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Ваши результаты очного тура")
        self.assertContains(resp, "Итоговый балл:")
        self.assertContains(resp, "18")  # 8 + 10 = 18
        self.assertContains(resp, "Задача 1.1")
        self.assertContains(resp, "Задача 1.2")
        # Задачи комплекта 2 не показываются участнику комплекта 1
        self.assertNotContains(resp, "Задача 2.1")
        self.assertNotContains(resp, "Задача 2.2")
        # У задач нет гиперссылок
        self.assertNotContains(resp, self.p1_1.get_absolute_url())
        self.assertNotContains(resp, self.p1_2.get_absolute_url())




class VenueStagePassedTests(TestCase):
    def setUp(self):
        self.season = Season.objects.create(
            year=2099, slug="s2099", title="Сезон 2099", is_active=True, is_published=True
        )
        self.admin = User.objects.create_user(
            "admin_user@b.ru", "pass12345", role=User.Role.ADMIN, is_staff=True, is_superuser=True
        )
        self.venue = Venue.objects.create(
            title="Тестовая площадка",
            region="Москва",
            city="Москва",
            address="ул. Примерная, 10",
            latitude=55.75,
            longitude=37.62,
            status=Venue.Status.APPROVED,
            stage_passed=False,
        )
        self.venue.seasons.add(self.season)

    def test_default_stage_passed_is_false(self):
        self.assertFalse(self.venue.stage_passed)

    def test_as_geojson_feature_contains_stage_passed(self):
        feature = self.venue.as_geojson_feature()
        self.assertIn("stage_passed", feature["properties"])
        self.assertFalse(feature["properties"]["stage_passed"])

        self.venue.stage_passed = True
        self.venue.save()
        feature_updated = self.venue.as_geojson_feature()
        self.assertTrue(feature_updated["properties"]["stage_passed"])

    def test_venues_geojson_endpoint_returns_stage_passed(self):
        res = self.client.get(reverse("venues:geojson"))
        self.assertEqual(res.status_code, 200)
        data = res.json()
        props = next(f["properties"] for f in data["features"] if f["properties"]["id"] == self.venue.pk)
        self.assertFalse(props["stage_passed"])

        self.venue.stage_passed = True
        self.venue.save()
        res = self.client.get(reverse("venues:geojson"))
        data = res.json()
        props = next(f["properties"] for f in data["features"] if f["properties"]["id"] == self.venue.pk)
        self.assertTrue(props["stage_passed"])

    def test_detail_page_shows_passed_badge_when_stage_passed(self):
        url = reverse("venues:detail", args=[self.venue.pk])
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertNotContains(res, "Этап прошёл")

        self.venue.stage_passed = True
        self.venue.save()
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Этап прошёл")

    def test_admin_changelist_shows_stage_passed(self):
        self.client.force_login(self.admin)
        url = reverse("admin:venues_venue_changelist")
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Этап прошёл")
        self.assertContains(res, f'name="form-0-stage_passed"')

    def test_admin_actions_mark_stage_passed(self):
        self.client.force_login(self.admin)
        url = reverse("admin:venues_venue_changelist")

        # Mark stage passed
        post_data = {
            "action": "mark_stage_passed",
            "_selected_action": [str(self.venue.pk)],
        }
        res = self.client.post(url, post_data)
        self.assertEqual(res.status_code, 302)
        self.venue.refresh_from_db()
        self.assertTrue(self.venue.stage_passed)

        # Mark stage not passed
        post_data["action"] = "mark_stage_not_passed"
        res = self.client.post(url, post_data)
        self.assertEqual(res.status_code, 302)
        self.venue.refresh_from_db()
        self.assertFalse(self.venue.stage_passed)


class VenueRegistrationLockTests(TestCase):
    """Тесты блокировки регистрации на площадку через админку и запрета записи."""

    def setUp(self):
        now = timezone.now()
        self.season = Season.objects.create(
            year=2099, slug="s2099", title="Сезон 2099", is_active=True, is_published=True
        )
        self.stage = Stage.objects.create(
            season=self.season,
            kind=Stage.Kind.OFFLINE,
            slug="offline-final",
            title="Очный тур",
            starts_at=now - timedelta(days=1),
            ends_at=now + timedelta(days=2),
            is_published=True,
            order=1,
        )
        self.admin = User.objects.create_user(
            "admin_lock@b.ru", "pass12345", role=User.Role.ADMIN, is_staff=True, is_superuser=True
        )
        self.student = User.objects.create_user(
            "student_lock@b.ru", "pass12345", role=User.Role.PARTICIPANT, email_confirmed=True
        )
        ParticipantProfile.objects.create(
            user=self.student,
            last_name="Иванов",
            first_name="Иван",
            birth_date=date(2010, 1, 1),
            birth_place="г. Москва",
            doc_type=DocumentType.PASSPORT_RF,
            doc_series="4510",
            doc_number="123456",
            doc_issued_at=date(2024, 1, 1),
            doc_issued_by="ГУ МВД",
            doc_division_code="770001",
            reg_address="г. Москва",
            grade=10,
            school="Школа №1",
            city="Москва",
            region="Москва",
            phone="+79991234567",
            telegram="@ivanov",
            acknowledged_at=now,
        )
        ConsentDocument.objects.create(
            user=self.student,
            file="consents/test.pdf",
            status=ConsentDocument.Status.APPROVED,
        )
        self.venue = Venue.objects.create(
            title="Тестовая площадка",
            region="Москва",
            city="Москва",
            address="ул. Примерная, 10",
            latitude=55.75,
            longitude=37.62,
            status=Venue.Status.APPROVED,
            registration_closed=False,
        )
        self.venue.seasons.add(self.season)

    def test_default_registration_closed_is_false(self):
        self.assertFalse(self.venue.registration_closed)

    def test_as_geojson_feature_contains_registration_closed(self):
        feature = self.venue.as_geojson_feature()
        self.assertIn("registration_closed", feature["properties"])
        self.assertFalse(feature["properties"]["registration_closed"])

        self.venue.registration_closed = True
        self.venue.save()
        feature_updated = self.venue.as_geojson_feature()
        self.assertTrue(feature_updated["properties"]["registration_closed"])

    def test_admin_changelist_shows_registration_closed_checkbox(self):
        self.client.force_login(self.admin)
        url = reverse("admin:venues_venue_changelist")
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'name="form-0-registration_closed"')
        self.assertNotContains(res, "toggle-registration")

    def test_admin_actions_lock_and_unlock_registration(self):
        self.client.force_login(self.admin)
        url = reverse("admin:venues_venue_changelist")

        # Lock
        post_data = {
            "action": "lock_registration",
            "_selected_action": [str(self.venue.pk)],
        }
        res = self.client.post(url, post_data)
        self.assertEqual(res.status_code, 302)
        self.venue.refresh_from_db()
        self.assertTrue(self.venue.registration_closed)

        # Unlock
        post_data["action"] = "unlock_registration"
        res = self.client.post(url, post_data)
        self.assertEqual(res.status_code, 302)
        self.venue.refresh_from_db()
        self.assertFalse(self.venue.registration_closed)

    def test_detail_page_and_booking_prevented_when_registration_closed(self):
        self.client.force_login(self.student)
        url = reverse("venues:detail", args=[self.venue.pk])

        # Registration open: booking button is shown
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Записаться")

        # Lock registration
        self.venue.registration_closed = True
        self.venue.save()

        # Detail page shows locked badge and no booking button
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Запись заблокирована")
        self.assertNotContains(res, ">Записаться<")

        # Attempt to book via POST is rejected
        book_url = reverse("venues:book", args=[self.venue.pk])
        res = self.client.post(book_url)
        self.assertEqual(res.status_code, 302)
        self.assertFalse(VenueBooking.objects.filter(venue=self.venue, registration__user=self.student).exists())


class VenueStagePassedCancelTests(TestCase):
    """Тесты запрета участникам отписываться от точек, когда очный этап прошёл."""

    def setUp(self):
        from apps.accounts.testing import make_eligible

        now = timezone.now()
        self.season = Season.objects.create(
            year=2099, slug="s2099", title="Сезон 2099", is_active=True, is_published=True
        )
        self.stage = Stage.objects.create(
            season=self.season,
            kind=Stage.Kind.OFFLINE,
            slug="offline-final",
            title="Очный тур",
            starts_at=now - timedelta(days=1),
            ends_at=now + timedelta(days=2),
            registration_opens_at=now - timedelta(days=5),
            is_published=True,
            order=1,
        )
        self.venue1 = Venue.objects.create(
            title="Площадка 1",
            region="Москва",
            city="Москва",
            address="ул. Первая, 1",
            latitude=55.75,
            longitude=37.62,
            status=Venue.Status.APPROVED,
            stage_passed=False,
        )
        self.venue1.seasons.add(self.season)

        self.venue2 = Venue.objects.create(
            title="Площадка 2",
            region="Москва",
            city="Москва",
            address="ул. Вторая, 2",
            latitude=55.76,
            longitude=37.63,
            status=Venue.Status.APPROVED,
            stage_passed=False,
        )
        self.venue2.seasons.add(self.season)

        self.make_eligible = make_eligible
        self.student = make_eligible(User.objects.create_user("student_cancel@b.ru", "pass12345", role=User.Role.PARTICIPANT))
        self.registration = Registration.objects.create(user=self.student, season=self.season)
        self.booking = VenueBooking.objects.create(
            registration=self.registration,
            stage=self.stage,
            venue=self.venue1,
            status=VenueBooking.Status.BOOKED,
        )

    def test_can_cancel_when_stage_has_not_passed(self):
        self.assertTrue(self.booking.can_cancel)
        self.client.force_login(self.student)
        url = reverse("venues:detail", args=[self.venue1.pk])
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Отменить запись")

        cancel_url = reverse("venues:cancel_booking", args=[self.venue1.pk])
        post_res = self.client.post(cancel_url)
        self.assertEqual(post_res.status_code, 302)
        self.booking.refresh_from_db()
        self.assertEqual(self.booking.status, VenueBooking.Status.CANCELLED)

    def test_cannot_cancel_when_venue_stage_passed(self):
        self.venue1.stage_passed = True
        self.venue1.save()
        self.booking.refresh_from_db()
        self.assertFalse(self.booking.can_cancel)

        self.client.force_login(self.student)
        url = reverse("venues:detail", args=[self.venue1.pk])
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertNotContains(res, 'class="link-button">Отменить запись</button>')
        self.assertContains(res, "Отменить запись нельзя: очный этап уже прошёл.")

        # Attempt to cancel via POST is rejected
        cancel_url = reverse("venues:cancel_booking", args=[self.venue1.pk])
        post_res = self.client.post(cancel_url)
        self.assertEqual(post_res.status_code, 302)
        self.booking.refresh_from_db()
        self.assertEqual(self.booking.status, VenueBooking.Status.BOOKED)

    def test_cannot_cancel_when_stage_deadline_is_past(self):
        self.stage.ends_at = timezone.now() - timedelta(hours=1)
        self.stage.save()
        self.booking.refresh_from_db()
        self.assertFalse(self.booking.can_cancel)

        self.client.force_login(self.student)
        url = reverse("venues:detail", args=[self.venue1.pk])
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertNotContains(res, 'class="link-button">Отменить запись</button>')
        self.assertContains(res, "Отменить запись нельзя: очный этап уже прошёл.")

        cancel_url = reverse("venues:cancel_booking", args=[self.venue1.pk])
        post_res = self.client.post(cancel_url)
        self.assertEqual(post_res.status_code, 302)
        self.booking.refresh_from_db()
        self.assertEqual(self.booking.status, VenueBooking.Status.BOOKED)

    def test_cannot_switch_venue_when_current_venue_stage_passed(self):
        self.venue1.stage_passed = True
        self.venue1.save()

        self.client.force_login(self.student)
        # Detail page for venue 2 should forbid switching
        url = reverse("venues:detail", args=[self.venue2.pk])
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertNotContains(res, "Перенести запись сюда")
        self.assertContains(res, "Сменить площадку нельзя.")
        self.assertContains(res, "где очный этап уже прошёл")

        # POST to book venue 2 is rejected
        book_url = reverse("venues:book", args=[self.venue2.pk])
        post_res = self.client.post(book_url)
        self.assertEqual(post_res.status_code, 302)
        self.booking.refresh_from_db()
        self.assertEqual(self.booking.venue, self.venue1)
        self.assertEqual(self.booking.status, VenueBooking.Status.BOOKED)

    def test_cannot_book_venue_when_target_venue_stage_passed(self):
        self.venue2.stage_passed = True
        self.venue2.save()

        other_student = self.make_eligible(User.objects.create_user("student_other@b.ru", "pass12345", role=User.Role.PARTICIPANT))
        self.client.force_login(other_student)

        url = reverse("venues:detail", args=[self.venue2.pk])
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertNotContains(res, ">Записаться<")
        self.assertContains(res, "Запись на эту площадку закрыта: очный этап уже прошёл.")

        book_url = reverse("venues:book", args=[self.venue2.pk])
        post_res = self.client.post(book_url)
        self.assertEqual(post_res.status_code, 302)
        self.assertFalse(VenueBooking.objects.filter(venue=self.venue2, registration__user=other_student).exists())

    def test_map_hides_cancel_link_when_venue_stage_passed(self):
        self.client.force_login(self.student)
        map_url = reverse("venues:map")

        # Initially stage not passed -> shows link with cancel mention
        res = self.client.get(map_url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Посмотреть площадку или отменить запись →")

        # Once venue stage_passed is marked -> does not suggest cancelling
        self.venue1.stage_passed = True
        self.venue1.save()
        res = self.client.get(map_url)
        self.assertEqual(res.status_code, 200)
        self.assertNotContains(res, "или отменить запись")
        self.assertContains(res, "Посмотреть площадку →")
